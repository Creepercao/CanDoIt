"""PPTX export for HTML slide presentations.

Primary path:
  HTML -> Playwright-rendered slide screenshots -> PPTX full-slide images.

This preserves CSS backgrounds, SVG/CSS graphics, layout, and visual styling
far better than trying to reconstruct a deck from extracted text. JavaScript
and CSS animations cannot be transferred into editable PowerPoint animation
objects, so ``mode="final"`` captures each slide after animations settle and
``mode="keyframes"`` captures several moments per slide as separate PPT pages.

Fallback path:
  HTML text extraction -> python-pptx text boxes.
Used when Playwright/Chromium is unavailable.
"""
from __future__ import annotations

import asyncio
import logging
import os
import re
import shutil
import uuid
from pathlib import Path
from typing import Optional

from bs4 import BeautifulSoup
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

logger = logging.getLogger("ppt_export")

OUTPUTS_DIR = Path(__file__).parent.parent.parent / "outputs"
SCREENSHOT_DIR = OUTPUTS_DIR / "pptx_frames"

_TITLE_RE = re.compile(r"<title>(.*?)</title>", re.IGNORECASE | re.DOTALL)
_HTML_RE = re.compile(r"<!DOCTYPE\s+html|<html[\s>]", re.IGNORECASE)


def _extract_title(html: str, fallback: str = "") -> str:
    if fallback:
        return fallback
    match = _TITLE_RE.search(html)
    if match:
        return BeautifulSoup(match.group(1), "html.parser").get_text(" ", strip=True)
    return "ppt-animation 演示文稿"


def _extract_slides_from_html(html: str) -> list[dict]:
    """Extract slide text for fallback export and slide count reporting."""
    soup = BeautifulSoup(html, "html.parser")
    for tag in soup(["script", "style", "head"]):
        tag.decompose()

    nodes = soup.select(".slide, .page, section, article")
    if not nodes:
        body_text = soup.get_text("\n", strip=True)
        return [{"title": "Slide 1", "content": body_text[:3000], "index": 0}] if body_text else []

    slides: list[dict] = []
    for node in nodes:
        text = node.get_text("\n", strip=True)
        if not text:
            continue
        title_node = node.find(["h1", "h2", "h3"])
        title = title_node.get_text(" ", strip=True) if title_node else f"Slide {len(slides) + 1}"
        if title and text.startswith(title):
            text = text[len(title):].strip()
        slides.append({"title": title, "content": text[:3000], "index": len(slides)})
    return slides


def _sanitize_capture_mode(mode: str) -> str:
    return "keyframes" if mode == "keyframes" else "final"


async def _capture_html_slides(
    html_content: str,
    *,
    mode: str = "final",
    frames_per_slide: int = 3,
    width: int = 1920,
    height: int = 1080,
) -> tuple[list[Path], int]:
    """Render HTML in Chromium and capture slide screenshots.

    Returns ``(image_paths, logical_slide_count)``. In keyframe mode,
    ``len(image_paths)`` may be ``logical_slide_count * frames_per_slide``.
    """
    try:
        from playwright.async_api import async_playwright
    except Exception as e:  # pragma: no cover - environment dependent
        raise RuntimeError("Playwright is not installed") from e

    mode = _sanitize_capture_mode(mode)
    frames_per_slide = max(2, min(int(frames_per_slide or 3), 6))
    SCREENSHOT_DIR.mkdir(parents=True, exist_ok=True)

    run_id = uuid.uuid4().hex[:12]
    html_path = SCREENSHOT_DIR / f"source_{run_id}.html"
    html_path.write_text(html_content, encoding="utf-8")

    image_paths: list[Path] = []
    logical_count = 0
    chromium_path = (
        os.environ.get("PLAYWRIGHT_CHROMIUM_EXECUTABLE")
        or shutil.which("chromium")
        or shutil.which("chromium-browser")
        or shutil.which("google-chrome")
    )

    async with async_playwright() as p:
        launch_kwargs = {"args": ["--no-sandbox"]}
        if chromium_path:
            launch_kwargs["executable_path"] = chromium_path
        browser = await p.chromium.launch(**launch_kwargs)
        page = await browser.new_page(viewport={"width": width, "height": height}, device_scale_factor=1)
        try:
            await page.goto(html_path.resolve().as_uri(), wait_until="load", timeout=30000)
            await page.wait_for_timeout(300)

            logical_count = await page.evaluate(
                """() => {
                    // Count slides by DOM presence, NOT by bounding-rect
                    // dimensions.  Most decks keep inactive slides at
                    // display:none, which makes getBoundingClientRect()
                    // return {width:0, height:0} and the old filter
                    // silently dropped them.
                    const slides = document.querySelectorAll('.slide, .page');
                    if (slides.length > 0) return slides.length;
                    // Fallback: count <section>/<article> inside the deck
                    const deck = document.querySelector('.deck, main, body');
                    const sections = deck
                        ? deck.querySelectorAll('section, article')
                        : document.querySelectorAll('section, article');
                    return sections.length || 1;
                }"""
            )

            # Force one logical slide visible at a time. This works for most
            # skill-generated decks and avoids needing to understand each deck's
            # custom navigation JavaScript.
            for slide_idx in range(logical_count):
                await page.evaluate(
                    """({idx}) => {
                        const slides = Array.from(document.querySelectorAll('.slide, .page, section, article'));
                        document.documentElement.style.width = '100%';
                        document.documentElement.style.height = '100%';
                        document.body.style.margin = '0';
                        document.body.style.width = '100vw';
                        document.body.style.height = '100vh';
                        document.body.style.overflow = 'hidden';

                        if (!slides.length) return;

                        slides.forEach((el, i) => {
                            el.classList.toggle('active', i === idx);
                            el.classList.toggle('current', i === idx);
                            el.style.display = i === idx ? 'block' : 'none';
                            el.style.visibility = i === idx ? 'visible' : 'hidden';
                            el.style.opacity = i === idx ? '1' : '0';
                            el.style.pointerEvents = i === idx ? 'auto' : 'none';
                            if (i === idx) {
                                el.style.position = 'fixed';
                                el.style.left = '0';
                                el.style.top = '0';
                                el.style.width = '100vw';
                                el.style.height = '100vh';
                                el.style.margin = '0';
                                el.style.transform = 'none';
                            }
                        });
                    }""",
                    {"idx": slide_idx},
                )

                waits = [2300] if mode == "final" else [
                    int(350 + (2200 * i / max(frames_per_slide - 1, 1)))
                    for i in range(frames_per_slide)
                ]

                for frame_idx, wait_ms in enumerate(waits):
                    await page.wait_for_timeout(wait_ms if frame_idx == 0 else max(wait_ms - waits[frame_idx - 1], 200))
                    suffix = f"s{slide_idx + 1:02d}"
                    if mode == "keyframes":
                        suffix += f"_f{frame_idx + 1:02d}"
                    img_path = SCREENSHOT_DIR / f"{run_id}_{suffix}.png"
                    await page.screenshot(path=str(img_path), full_page=False)
                    image_paths.append(img_path)
        finally:
            await browser.close()

    return image_paths, logical_count


def _images_to_pptx(image_paths: list[Path], output_path: str) -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    for image_path in image_paths:
        slide = prs.slides.add_slide(blank_layout)
        slide.shapes.add_picture(
            str(image_path),
            0,
            0,
            width=prs.slide_width,
            height=prs.slide_height,
        )

    prs.save(output_path)


def _fallback_text_to_pptx(html_content: str, output_path: str, title: str) -> int:
    """Old text extraction fallback, retained for environments without Chromium."""
    slides = _extract_slides_from_html(html_content)
    if not slides:
        slides = [{"title": title, "content": "No slide content found.", "index": 0}]

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    for s in slides:
        slide_obj = prs.slides.add_slide(blank_layout)
        left, top = Inches(0.8), Inches(0.4)
        width, title_h = Inches(11.7), Inches(1.0)

        title_box = slide_obj.shapes.add_textbox(left, top, width, title_h)
        p = title_box.text_frame.paragraphs[0]
        p.text = s["title"]
        p.font.size = Pt(32)
        p.font.bold = True
        p.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)
        p.alignment = PP_ALIGN.LEFT

        content_box = slide_obj.shapes.add_textbox(left, Inches(1.6), width, Inches(5.5))
        frame = content_box.text_frame
        frame.word_wrap = True
        paragraphs = [p for p in s["content"].split("\n") if p.strip()] or [s["content"]]
        for j, para_text in enumerate(paragraphs[:15]):
            para = frame.paragraphs[0] if j == 0 else frame.add_paragraph()
            para.text = para_text[:500]
            para.font.size = Pt(18)
            para.font.color.rgb = RGBColor(0x33, 0x33, 0x44)
            para.space_after = Pt(8)
            para.alignment = PP_ALIGN.LEFT

    prs.save(output_path)
    return len(slides)


async def html_to_pptx_async(
    html_content: str,
    output_path: Optional[str] = None,
    title: str = "Presentation",
    mode: str = "final",
    frames_per_slide: int = 3,
) -> dict:
    if output_path is None:
        output_path = str(OUTPUTS_DIR / f"export_{uuid.uuid4().hex[:12]}.pptx")

    if not _HTML_RE.search(html_content[:1000]):
        raise ValueError("Content does not appear to be HTML")

    try:
        image_paths, logical_slides = await _capture_html_slides(
            html_content,
            mode=mode,
            frames_per_slide=frames_per_slide,
        )
        if not image_paths:
            raise RuntimeError("No screenshots captured")
        _images_to_pptx(image_paths, output_path)
        logger.info(
            "PPTX exported via screenshots: %s (%s slide(s), %s image page(s), mode=%s)",
            output_path,
            logical_slides,
            len(image_paths),
            mode,
        )
        return {
            "file_path": output_path,
            "slides": logical_slides,
            "pages": len(image_paths),
            "mode": _sanitize_capture_mode(mode),
            "rendered": "screenshot",
        }
    except Exception as e:
        logger.warning("Screenshot PPTX export failed, falling back to text export: %s", e)
        slide_count = _fallback_text_to_pptx(html_content, output_path, title)
        return {
            "file_path": output_path,
            "slides": slide_count,
            "pages": slide_count,
            "mode": "fallback-text",
            "rendered": "text",
            "warning": str(e),
        }


async def export_skill_to_pptx_async(
    html_content: str,
    skill_name: str = "ppt-animation",
    title: str = "",
    mode: str = "final",
    frames_per_slide: int = 3,
) -> dict:
    title = _extract_title(html_content, title)
    file_id = uuid.uuid4().hex[:12]
    output_path = str(OUTPUTS_DIR / f"{skill_name}_{file_id}.pptx")

    result = await html_to_pptx_async(
        html_content,
        output_path=output_path,
        title=title,
        mode=mode,
        frames_per_slide=frames_per_slide,
    )
    file_path = Path(result["file_path"])
    return {
        "file_url": f"/outputs/{file_path.name}",
        "download_url": f"/api/skills/ppt-animation/export-pptx/{file_path.name}",
        "file_path": str(file_path),
        "slides": result["slides"],
        "pages": result["pages"],
        "title": title,
        "mode": result["mode"],
        "rendered": result["rendered"],
        "warning": result.get("warning", ""),
    }


def export_skill_to_pptx(
    html_content: str,
    skill_name: str = "ppt-animation",
    title: str = "",
    mode: str = "final",
    frames_per_slide: int = 3,
) -> dict:
    """Synchronous wrapper for scripts/tests outside FastAPI."""
    return asyncio.run(export_skill_to_pptx_async(
        html_content,
        skill_name=skill_name,
        title=title,
        mode=mode,
        frames_per_slide=frames_per_slide,
    ))
